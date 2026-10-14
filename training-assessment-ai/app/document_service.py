from io import BytesIO
from pathlib import Path
from docx import Document
from pptx import Presentation
from pypdf import PdfReader

SUPPORTED_EXTENSIONS = {".pdf", ".pptx", ".docx", ".txt", ".md"}


class DocumentError(ValueError):
    pass


def extract_text(filename: str, content: bytes, max_chars: int = 100000) -> str:
    ext = Path(filename or "").suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise DocumentError("Unsupported file type. Upload a PDF, PPTX, DOCX, TXT, or Markdown file.")
    if not content:
        raise DocumentError("The uploaded file is empty.")
    try:
        if ext == ".pdf":
            reader = PdfReader(BytesIO(content), strict=False)
            pages = [(page.extract_text() or "").strip() for page in reader.pages]
            text = "\n\n".join(f"[Page {i + 1}]\n{page}" for i, page in enumerate(pages) if page)
            if not text.strip():
                raise DocumentError("No selectable text was found in this PDF. It may be scanned; OCR is not implemented, so upload an OCR-processed PDF or text document.")
        elif ext == ".docx":
            doc = Document(BytesIO(content))
            text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
            for table in doc.tables:
                text += "\n" + "\n".join(" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows)
        elif ext == ".pptx":
            prs = Presentation(BytesIO(content))
            chunks = []
            for i, slide in enumerate(prs.slides, start=1):
                lines = [shape.text.strip() for shape in slide.shapes if getattr(shape, "has_text_frame", False) and shape.text.strip()]
                if lines:
                    chunks.append(f"[Slide {i}]\n" + "\n".join(lines))
            text = "\n\n".join(chunks)
        else:
            text = content.decode("utf-8-sig")
    except DocumentError:
        raise
    except Exception as exc:
        raise DocumentError("The document could not be read. Check that it is a valid, uncorrupted file.") from exc
    text = text.replace("\x00", "").strip()
    if not text:
        raise DocumentError("No readable text was found in this document.")
    if len(text) > max_chars:
        raise DocumentError(f"Extracted text is too long ({len(text):,} characters). Limit it to {max_chars:,} characters or split the material into smaller files.")
    return text
